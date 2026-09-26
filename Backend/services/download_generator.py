from io import BytesIO

from docx import Document
from docx.shared import Pt

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics


def create_docx(text, filename, language):
    """
    Create a DOCX document from translated text.
    """

    document = Document()

    # Title
    title = document.add_heading(
        "Learn A2Z",
        level=0
    )

    subtitle = document.add_paragraph(
        f"Translated Learning Document — {language}"
    )

    subtitle.runs[0].font.size = Pt(11)

    document.add_paragraph("")

    # Document content
    paragraphs = text.split("\n")

    for paragraph_text in paragraphs:

        paragraph_text = paragraph_text.strip()

        if not paragraph_text:
            document.add_paragraph("")
            continue

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            paragraph_text
        )

        run.font.size = Pt(11)

    # Save into memory
    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output


def create_pdf(text, filename, language):
    """
    Create a PDF document from translated text.
    """

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=50,
        bottomMargin=50,
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    body_style = styles["BodyText"]

    body_style.fontSize = 11
    body_style.leading = 17
    body_style.alignment = TA_LEFT

    story = []

    # Title
    story.append(
        Paragraph(
            "Learn A2Z",
            title_style
        )
    )

    story.append(
        Spacer(1, 8)
    )

    story.append(
        Paragraph(
            f"Translated Learning Document — {language}",
            styles["Normal"]
        )
    )

    story.append(
        Spacer(1, 20)
    )

    # Content
    paragraphs = text.split("\n")

    for paragraph_text in paragraphs:

        paragraph_text = paragraph_text.strip()

        if not paragraph_text:
            story.append(
                Spacer(1, 8)
            )
            continue

        # Escape special XML characters
        paragraph_text = (
            paragraph_text
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        story.append(
            Paragraph(
                paragraph_text,
                body_style
            )
        )

        story.append(
            Spacer(1, 8)
        )

    document.build(story)

    output.seek(0)

    return output